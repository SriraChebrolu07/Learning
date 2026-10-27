# import pdfplumber
# from docx import Document

# def pdf_to_word(pdf_path, docx_path):
#     # Create a new Document
#     doc = Document()

#     # Open the PDF file
#     with pdfplumber.open(pdf_path) as pdf:
#         # Iterate through all the pages
#         for page in pdf.pages:
#             # Extract text from the page
#             text = page.extract_text()
#             if text:
#                 # Add text to the Word document
#                 doc.add_paragraph(text)

#     # Save the Word document
#     doc.save(docx_path)
#     print(f"Successfully converted {pdf_path} to {docx_path}")

# # Example usage
# if __name__ == '__main__':
#     pdf_path = 'aws.pdf'  # Replace with your PDF file path
#     docx_path = 'output.docx'  # Replace with desired output Word file path
#     pdf_to_word(pdf_path, docx_path)




# import PyPDF2
# from docx import Document
# with open('aws pdf','r') as file:
#     pdf=PyPDF2.PdfFileReader(file)
#     document=Document()
#     for page in range(pdf.getNumPages()):
#         text = pdf.getPage(page).extractText()
#         document.add_paragraph(text)
# document.save('vpc.docx')






import PyPDF2
from docx import Document

# Function to convert PDF to Word
def pdf_to_word(pdf_path, docx_path):
    # Open the PDF file
    with open(pdf_path, 'rb') as file:
        pdf = PyPDF2.PdfFileReader(file)
        document = Document()
        
        # Iterate through all the pages
        for page in range(pdf.getNumPages()):
            # Extract text from each page
            text = pdf.getPage(page).extract_text()
            # Add text to the Word document
            document.add_paragraph(text)
            
    # Save the Word document
    document.save(docx_path)
    print(f"Successfully converted {pdf_path} to {docx_path}")

# Example usage
if __name__ == '__main__':
    pdf_path = 'aws.pdf'  # Replace with your PDF file path
    docx_path = 'output.docx'  # Replace with desired output Word file path
    pdf_to_word(pdf_path, docx_path)
